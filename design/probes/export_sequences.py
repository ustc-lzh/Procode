#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Write sequence tables as colored XLSX files with a standard-library fallback."""

from __future__ import annotations

import os
import zipfile
from typing import List, Optional, Sequence
from xml.sax.saxutils import escape

__all__ = ["write_colored_xlsx", "HAVE_OPENPYXL"]

try:
    import openpyxl                                          # noqa: F401
    HAVE_OPENPYXL = True
except ImportError:
    HAVE_OPENPYXL = False


def write_colored_xlsx(path: str, header: Sequence[str], rows: Sequence[Sequence[object]],
                       fills: Optional[dict] = None) -> None:
    """Write an XLSX file.

    - header/rows: column headers and data
    - fills: {(row_index_0based, col_index_0based): "RRGGBB"}; only cells to color
    """
    os.makedirs(os.path.dirname(os.path.abspath(path)) or ".", exist_ok=True)
    fills = fills or {}
    if HAVE_OPENPYXL:
        _write_openpyxl(path, header, rows, fills)
    else:
        _write_minimal(path, header, rows, fills)


def _write_openpyxl(path, header, rows, fills) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import PatternFill

    wb = Workbook()
    ws = wb.active
    ws.append(list(header))
    for row in rows:
        ws.append(list(row))
    cache = {}
    for (r, c), color in fills.items():
        fill = cache.get(color)
        if fill is None:
            fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
            cache[color] = fill
        ws.cell(row=r + 2, column=c + 1).fill = fill      # +2: 1-based; row 1 is the header
    wb.save(path)


# --------------------------------------------------------------------------- #
# Minimal XLSX (no third-party dependencies)
# --------------------------------------------------------------------------- #
_CONTENT_TYPES = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
<Default Extension="xml" ContentType="application/xml"/>
<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>
<Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>
<Override PartName="/xl/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml"/>
</Types>"""

_ROOT_RELS = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/>
</Relationships>"""

_WORKBOOK = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
<sheets><sheet name="Sheet1" sheetId="1" r:id="rId1"/></sheets>
</workbook>"""

_WORKBOOK_RELS = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/>
<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/>
</Relationships>"""


def _col_letter(index: int) -> str:
    letters = ""
    index += 1
    while index:
        index, rem = divmod(index - 1, 26)
        letters = chr(65 + rem) + letters
    return letters


def _styles_xml(colors: Sequence[str]) -> str:
    fills = ['<fill><patternFill patternType="none"/></fill>',
             '<fill><patternFill patternType="gray125"/></fill>']
    for color in colors:
        fills.append('<fill><patternFill patternType="solid">'
                     '<fgColor rgb="FF%s"/><bgColor indexed="64"/>'
                     '</patternFill></fill>' % color)
    xfs = ['<xf numFmtId="0" fontId="0" fillId="0" borderId="0" xfId="0"/>']
    for i in range(len(colors)):
        xfs.append('<xf numFmtId="0" fontId="0" fillId="%d" borderId="0" xfId="0" applyFill="1"/>'
                   % (i + 2))
    return ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<styleSheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
            '<fonts count="1"><font><sz val="11"/><name val="Calibri"/></font></fonts>'
            '<fills count="%d">%s</fills>'
            '<borders count="1"><border/></borders>'
            '<cellStyleXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0"/></cellStyleXfs>'
            '<cellXfs count="%d">%s</cellXfs>'
            '</styleSheet>' % (len(fills), "".join(fills), len(xfs), "".join(xfs)))


def _write_minimal(path, header, rows, fills) -> None:
    colors: List[str] = sorted(set(fills.values()))
    style_of = {color: i + 1 for i, color in enumerate(colors)}   # cellXfs index

    parts = ['<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
             '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
             '<sheetData>']
    parts.append('<row r="1">')
    for c, value in enumerate(header):
        parts.append('<c r="%s1" t="inlineStr"><is><t xml:space="preserve">%s</t></is></c>'
                     % (_col_letter(c), escape(str(value))))
    parts.append('</row>')

    for r, row in enumerate(rows):
        parts.append('<row r="%d">' % (r + 2))
        for c, value in enumerate(row):
            style = fills.get((r, c))
            attr = ' s="%d"' % style_of[style] if style else ""
            parts.append('<c r="%s%d"%s t="inlineStr"><is><t xml:space="preserve">%s</t></is></c>'
                         % (_col_letter(c), r + 2, attr, escape("" if value is None else str(value))))
        parts.append('</row>')
    parts.append('</sheetData></worksheet>')

    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("[Content_Types].xml", _CONTENT_TYPES)
        zf.writestr("_rels/.rels", _ROOT_RELS)
        zf.writestr("xl/workbook.xml", _WORKBOOK)
        zf.writestr("xl/_rels/workbook.xml.rels", _WORKBOOK_RELS)
        zf.writestr("xl/styles.xml", _styles_xml(colors))
        zf.writestr("xl/worksheets/sheet1.xml", "".join(parts))
